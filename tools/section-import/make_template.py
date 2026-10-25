"""
Builds the TermPoint section import template workbook.

Writes src/TermPoint/Assets/Templates/TermPointSectionsTemplate.xlsx with two sheets:
  * Sections     - the 18 template columns, a few sample rows that show every pattern, and Text
                   formatting on the code and time columns so Excel keeps leading zeros.
  * Instructions - built from src/TermPoint/Assets/Templates/SectionImportInstructions.txt. That one
                   file is also what the app's Import Sections flyout shows and what its "Save instructions..."
                   button saves as a web page, so edit it there and re-run this script.

Run from anywhere:   python tools/section-import/make_template.py
Needs openpyxl:      python -m pip install --user openpyxl

The COLUMNS list below must match SectionImportColumns.All in the app, spelled and ordered
identically. A C# test (TemplateWorkbookTests) fails if the shipped workbook drifts from it.
"""

import datetime
import io
import math
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------------------
# Locations
# --------------------------------------------------------------------------------------

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent.parent
INSTRUCTIONS_PATH = REPO_ROOT / "src" / "TermPoint" / "Assets" / "Templates" / "SectionImportInstructions.txt"
OUTPUT_PATH = REPO_ROOT / "src" / "TermPoint" / "Assets" / "Templates" / "TermPointSectionsTemplate.xlsx"

# --------------------------------------------------------------------------------------
# The Sections sheet
# --------------------------------------------------------------------------------------

# (header, column width). Order and spelling must match SectionImportColumns.All.
COLUMNS = [
    ("Academic Year", 14),
    ("Term", 9),
    ("Subject", 9),
    ("Course Number", 15),
    ("Course Title", 28),
    ("Section Code", 13),
    ("Section Type", 14),
    ("Instructor(s)", 26),
    ("Tags", 20),
    ("Capacity", 10),
    ("Meeting Days", 14),
    ("Begin", 8),
    ("End", 8),
    ("Frequency", 11),
    ("Meeting Type", 14),
    ("Building", 10),
    ("Room", 8),
    ("Campus", 10),
]

# Columns whose cells are formatted as Text so Excel does not turn 001 into 1 or 0930 into 930.
TEXT_COLUMNS = ["Course Number", "Section Code", "Begin", "End"]

# Rows 2..LAST_FORMATTED_ROW carry the Text format, so rows the user adds later behave too.
LAST_FORMATTED_ROW = 2000

# Sample rows, written as strings in COLUMNS order. They must plan with no problems
# (TemplateWorkbookTests checks this against an empty database).
AY, TERM = "2026-2027", "Fall"
SAMPLE_ROWS = [
    # Two instructors, three days, a room, two tags, a capacity, a campus.
    [AY, TERM, "CHEM", "101", "Introductory Chemistry", "A", "Lecture", "Smith, John; Lee, Ann",
     "Science; First Year", "60", "M,W,F", "0900", "0950", "", "Class", "SCI", "210", "Main"],
    # A lab that meets on odd weeks.
    [AY, TERM, "CHEM", "101", "Introductory Chemistry", "L1", "Lab", "Lee, Ann",
     "Science", "24", "R", "1300", "1550", "odd", "Lab", "SCI", "105", "Main"],
    # One section on two rows: different times, different rooms. The second row leaves the
    # section-level cells blank.
    [AY, TERM, "ENGL", "110", "Academic Writing", "A", "Seminar", "Brown, Maria",
     "Writing", "25", "M,W", "1000", "1120", "", "Class", "ARTS", "214", "Main"],
    [AY, TERM, "ENGL", "110", "", "A", "", "",
     "", "", "F", "1300", "1350", "", "Class", "ARTS", "118", ""],
    # An online section: no meeting days or times, and a leading-zero section code.
    [AY, TERM, "HIST", "200", "Topics in History", "001", "Online", "Garcia, Luis",
     "", "30", "", "", "", "", "", "", "", ""],
]

# --------------------------------------------------------------------------------------
# Styling
# --------------------------------------------------------------------------------------

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")
HEADING_FONT = Font(bold=True, size=13)
INSTRUCTIONS_COLUMN_WIDTH = 110

# Fixed workbook dates, so rebuilding unchanged inputs gives the same file.
FIXED_DATE = datetime.datetime(2026, 10, 25, 0, 0, 0)
FIXED_ZIP_TIME = (2026, 10, 25, 0, 0, 0)


def build_sections_sheet(ws):
    """Fills the Sections sheet: headers, Text formatting, sample rows, widths, frozen header."""
    ws.title = "Sections"

    for col, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    # Text format on the code and time columns for a generous block of rows (header included).
    names = [h for h, _ in COLUMNS]
    for header in TEXT_COLUMNS:
        col = names.index(header) + 1
        for row in range(1, LAST_FORMATTED_ROW + 1):
            ws.cell(row=row, column=col).number_format = "@"

    # Sample rows: every value a string (so 001 and 0930 are stored as text), except Capacity,
    # which is a real number so Excel doesn't flag it as "a number stored as text".
    capacity_col = names.index("Capacity") + 1
    for r, values in enumerate(SAMPLE_ROWS, start=2):
        for col, value in enumerate(values, start=1):
            if value == "":
                continue
            ws.cell(row=r, column=col, value=int(value) if col == capacity_col else value)


def parse_instructions(text):
    """
    Turns SectionImportInstructions.txt into a list of (kind, text) blocks.

    Markup: '# ' starts a heading; '- ' starts a bullet (one line each); blank lines separate
    paragraphs; consecutive plain lines are joined into one paragraph.
    """
    blocks = []
    paragraph = []

    def flush():
        if paragraph:
            blocks.append(("paragraph", " ".join(paragraph)))
            paragraph.clear()

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            flush()
        elif line.startswith("# "):
            flush()
            blocks.append(("heading", line[2:].strip()))
        elif line.startswith("- "):
            flush()
            blocks.append(("bullet", line[2:].strip()))
        else:
            paragraph.append(line)
    flush()
    return blocks


def estimated_height(text, bold_size=11):
    """A generous row height (points) for wrapped text in the instructions column."""
    chars_per_line = INSTRUCTIONS_COLUMN_WIDTH * 0.85
    lines = max(1, math.ceil(len(text) / chars_per_line))
    return lines * (bold_size + 4.5) + 3


def build_instructions_sheet(ws, text):
    """Fills the Instructions sheet: one block per row in column A, wide and wrapped."""
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = INSTRUCTIONS_COLUMN_WIDTH
    ws.sheet_view.showGridLines = False

    row = 1
    first = True
    for kind, content in parse_instructions(text):
        if kind == "heading" and not first:
            row += 1  # a blank row of breathing space before each heading
        first = False

        cell = ws.cell(row=row, column=1)
        if kind == "heading":
            cell.value = content
            cell.font = HEADING_FONT
            ws.row_dimensions[row].height = 21
        elif kind == "bullet":
            cell.value = "•  " + content
            cell.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
            ws.row_dimensions[row].height = estimated_height(content)
        else:
            cell.value = content
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = estimated_height(content)
        row += 1


def make_reproducible(xlsx_bytes):
    """
    Rewrites the workbook's zip so the same inputs give the same bytes: every entry gets a fixed
    timestamp, and the 'modified' date in docProps/core.xml (which openpyxl sets to 'now' while
    saving) is replaced with the fixed date.
    """
    stamp = FIXED_DATE.strftime("%Y-%m-%dT%H:%M:%SZ")
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(xlsx_bytes)) as source, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + stamp.encode() + rb"\g<2>", data)
            new_info = zipfile.ZipInfo(info.filename, FIXED_ZIP_TIME)
            new_info.compress_type = zipfile.ZIP_DEFLATED
            new_info.external_attr = info.external_attr
            target.writestr(new_info, data)
    return out.getvalue()


def main():
    """Builds the workbook and writes it to OUTPUT_PATH."""
    instructions = INSTRUCTIONS_PATH.read_text(encoding="utf-8")

    wb = Workbook()
    build_sections_sheet(wb.active)
    build_instructions_sheet(wb.create_sheet(), instructions)

    wb.properties.creator = "TermPoint"
    wb.properties.lastModifiedBy = "TermPoint"
    wb.properties.title = "TermPoint Sections Template"
    wb.properties.created = FIXED_DATE
    wb.properties.modified = FIXED_DATE

    buffer = io.BytesIO()
    wb.save(buffer)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_bytes(make_reproducible(buffer.getvalue()))
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    sys.exit(main())
